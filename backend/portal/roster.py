"""Bounded CSV/XLSX parsing and row-by-row roster import."""

import csv
import io
import uuid
import zipfile
from pathlib import Path
from xml.etree.ElementTree import ParseError

from defusedxml.common import DefusedXmlException
from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import IntegrityError, transaction
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError as APIValidationError

from .models import AuditLog, Batch, StudentBatch, User

HEADERS = ("name", "email", "usn", "batch")


def parse_roster(upload):
    if upload.size > 2 * 1024 * 1024:
        raise APIValidationError("Roster must be at most 2 MiB.")
    suffix = Path(upload.name).suffix.lower()
    data = upload.read()
    try:
        if suffix == ".csv":
            values = list(csv.reader(io.StringIO(data.decode("utf-8-sig")), strict=True))
        elif suffix == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 10 * 1024 * 1024:
                    raise APIValidationError("Expanded workbook exceeds 10 MiB.")
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            try:
                sheet = workbook.active
                if sheet.max_column != 4:
                    raise APIValidationError("Workbook must have exactly four columns.")
                if sheet.max_row and sheet.max_row > settings.MAX_ROSTER_ROWS + 1:
                    raise APIValidationError("Roster exceeds the 1000-row limit.")
                values = []
                for cells in sheet.iter_rows():
                    if any(cell.data_type == "f" for cell in cells):
                        raise APIValidationError("Workbook formulas are not accepted; export values only.")
                    values.append([cell.value for cell in cells])
                    if len(values) > settings.MAX_ROSTER_ROWS + 1:
                        raise APIValidationError("Roster exceeds the 1000-row limit.")
            finally:
                workbook.close()
        else:
            raise APIValidationError("Use a UTF-8 .csv or .xlsx file.")
    except (
        UnicodeError,
        csv.Error,
        zipfile.BadZipFile,
        KeyError,
        ValueError,
        ParseError,
        DefusedXmlException,
    ) as exc:
        raise APIValidationError("Cannot read roster. Check its file format.") from exc
    if not values or tuple(str(v).strip().lower() for v in values[0]) != HEADERS:
        raise APIValidationError("Headers must be exactly: name,email,usn,batch.")
    if not 1 <= len(values) - 1 <= settings.MAX_ROSTER_ROWS:
        raise APIValidationError("Provide 1–1000 data rows.")
    return values[1:]


def import_roster(upload, actor):
    rows = parse_roster(upload)
    created, errors = 0, []
    for number, row in enumerate(rows, start=2):
        try:
            if len(row) != 4 or any(v is None for v in row):
                raise ValueError("Provide all four columns.")
            name, email, usn, batch_names = [str(v).strip() for v in row]
            email, usn = email.lower(), usn.upper()
            if not name or len(name) > 100 or not usn or len(usn) > 20 or len(email) > 150:
                raise ValueError("Invalid name, email length or USN.")
            validate_email(email)
            names = list(dict.fromkeys(x.strip() for x in batch_names.split("|") if x.strip()))
            batches = list(Batch.objects.filter(name__in=names))
            if not names or len(batches) != len(names):
                raise ValueError("Unknown batch. Create each named batch first.")
            with transaction.atomic():
                if User.objects.filter(email__iexact=email).exists():
                    raise ValueError("Duplicate email.")
                if User.objects.filter(usn=usn).exists():
                    raise ValueError("Duplicate USN.")
                student = User(
                    username=uuid.uuid4().hex,
                    name=name,
                    email=email,
                    usn=usn,
                    role="STUDENT",
                    must_reset_password=True,
                )
                student.set_unusable_password()
                student.save()
                StudentBatch.objects.bulk_create(
                    [StudentBatch(student=student, batch=batch) for batch in batches]
                )
                AuditLog.objects.create(actor=actor, action="STUDENT_IMPORTED", target=str(student.id))
            created += 1
        except (ValueError, ValidationError, IntegrityError) as exc:
            reason = "Duplicate or conflicting roster data." if isinstance(exc, IntegrityError) else str(exc)
            errors.append({"row": number, "reason": reason})
    return {"created": created, "errors": errors}
